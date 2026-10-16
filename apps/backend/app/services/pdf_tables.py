import io
import re

import pandas as pd
import pdfplumber
from openpyxl import Workbook

from app.services.perf_analyzer import quote

MAX_PAGES = 30
MAX_TABLES = 20
MAX_ROWS = 2000
MAX_COLS = 50
MAX_CELLS = 200_000


class PdfInputError(ValueError):
    pass


# ---------- extraction ----------

def _clean(cell) -> str:
    return "" if cell is None else re.sub(r"\s+", " ", str(cell)).strip()


def normalize_table(raw: list[list]) -> list[list[str]]:
    rows = [[_clean(c) for c in r] for r in raw if r]
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    rows = [r for r in rows if any(r)]  # drop empty rows
    keep = [i for i in range(width) if any(r[i] for r in rows)]  # drop empty columns
    return [[r[i] for i in keep] for r in rows]


_NUMERIC = re.compile(r"^[+-]?[\d,]*\.?\d+%?$")


def guess_header(rows: list[list[str]]) -> bool:
    if len(rows) < 2:
        return False
    first = rows[0]
    if any(not c for c in first) or len({c.lower() for c in first}) != len(first):
        return False
    return not all(_NUMERIC.match(c) for c in first)


def _collect(pages, settings: dict | None, method: str, notes: list[str]) -> list[dict]:
    tables: list[dict] = []
    for page in pages:
        for raw in page.extract_tables(table_settings=settings):
            rows = normalize_table(raw)
            if len(rows) < 2 or len(rows[0]) < 2:
                continue
            truncated = len(rows) > MAX_ROWS or len(rows[0]) > MAX_COLS
            rows = [r[:MAX_COLS] for r in rows[:MAX_ROWS]]
            tables.append({
                "id": f"p{page.page_number}-t{len(tables) + 1}", "page": page.page_number,
                "method": method, "rows": rows, "header_guess": guess_header(rows), "truncated": truncated,
            })
            if len(tables) >= MAX_TABLES:
                notes.append(f"Only the first {MAX_TABLES} tables are shown.")
                return tables
    return tables


def extract_tables(data: bytes) -> dict:
    notes: list[str] = []
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:
        if "password" in exc.__class__.__name__.lower():
            raise PdfInputError("The PDF is password protected")
        raise PdfInputError("Could not read this PDF")
    with pdf:
        total = len(pdf.pages)
        pages = pdf.pages[:MAX_PAGES]
        if total > MAX_PAGES:
            notes.append(f"Only the first {MAX_PAGES} of {total} pages were scanned.")
        if not any(p.chars for p in pages):
            raise PdfInputError("This looks like a scanned PDF (no selectable text). OCR is not supported. Use a text-based PDF, or an Excel/CSV file.")
        try:
            tables = _collect(pages, None, "lines", notes)
            if not tables:
                tables = _collect(pages, {"vertical_strategy": "text", "horizontal_strategy": "text"}, "text-alignment", notes)
                if tables:
                    notes.append("No ruled tables were found. These tables were guessed from text alignment and may be wrong. Check them carefully.")
        except Exception:
            raise PdfInputError("Could not extract tables from this PDF")
    if not tables:
        notes.append("No tables were detected in this PDF.")
    if any(t["truncated"] for t in tables):
        notes.append(f"Large tables were cut to {MAX_ROWS} rows and {MAX_COLS} columns.")
    return {"pages_total": total, "pages_scanned": len(pages), "tables": tables, "notes": notes}


# ---------- table -> frame -> Excel ----------

def _unique_headers(headers: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for i, h in enumerate(headers):
        base = h.strip() or f"column_{i + 1}"
        name, n = base, 2
        while name.lower() in seen:
            name = f"{base}_{n}"
            n += 1
        seen.add(name.lower())
        out.append(name)
    return out


def build_frame(headers: list[str], rows: list[list[str]]) -> pd.DataFrame:
    n = len(headers)
    if n == 0:
        raise PdfInputError("The table has no columns")
    if not rows:
        raise PdfInputError("The table has no data rows")
    if n * len(rows) > MAX_CELLS:
        raise PdfInputError("The table is too large")
    if any(len(c) > 5000 for r in rows for c in r):
        raise PdfInputError("A cell is longer than 5000 characters")
    fixed = [(list(r) + [""] * n)[:n] for r in rows]
    return pd.DataFrame(fixed, columns=_unique_headers(headers), dtype=str)


def build_xlsx(df: pd.DataFrame, types: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for j, name in enumerate(df.columns, start=1):
        ws.cell(row=1, column=j, value=str(name))
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, (v, t) in enumerate(zip(row, types), start=1):
            v = v.strip() if isinstance(v, str) else v
            if v == "" or v is None:
                continue  # truly empty cell
            if t in ("integer", "bigint"):
                value = int(v)
            elif t == "float":
                value = float(v)
            else:
                value = v
            cell = ws.cell(row=i, column=j, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # keep text that starts with "=" as text, never as a formula
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- starter queries ----------

def starter_queries(db_kind: str, schema: str, table: str, columns: list[dict], row_count: int) -> list[dict]:
    def q(n: str) -> str:
        return quote(db_kind, n)

    tbl = f"{q(schema)}.{q(table)}" if db_kind == "postgresql" else q(table)
    out = [
        {"title": "Preview the first rows", "sql": f"SELECT * FROM {tbl} LIMIT 100"},
        {"title": "Count the rows", "sql": f"SELECT COUNT(*) AS total_rows FROM {tbl}"},
    ]
    nums = [c for c in columns if c["inferred_type"] in ("integer", "bigint", "float")]
    cats = [c for c in columns if c["inferred_type"] in ("varchar", "text")
            and 2 <= c["distinct_count"] <= min(50, max(2, row_count - 1))]
    dates = [c for c in columns if c["inferred_type"] in ("date", "timestamp")]

    if nums:
        n = q(nums[0]["name"])
        out.append({"title": f"Summary of {nums[0]['name']}",
                    "sql": f"SELECT MIN({n}) AS min_value, MAX({n}) AS max_value, AVG({n}) AS avg_value FROM {tbl}"})
    if cats:
        c = q(cats[0]["name"])
        out.append({"title": f"Rows per {cats[0]['name']}",
                    "sql": f"SELECT {c}, COUNT(*) AS total FROM {tbl} GROUP BY {c} ORDER BY total DESC LIMIT 20"})
        if nums:
            n = q(nums[0]["name"])
            out.append({"title": f"Average {nums[0]['name']} per {cats[0]['name']}",
                        "sql": f"SELECT {c}, AVG({n}) AS avg_value FROM {tbl} GROUP BY {c} ORDER BY avg_value DESC LIMIT 20"})
    if dates:
        d = q(dates[0]["name"])
        out.append({"title": f"Range of {dates[0]['name']}", "sql": f"SELECT MIN({d}) AS first_value, MAX({d}) AS last_value FROM {tbl}"})
    missing = next((c for c in columns if c["null_count"] > 0), None)
    if missing:
        out.append({"title": f"Rows with no {missing['name']}",
                    "sql": f"SELECT COUNT(*) AS missing_rows FROM {tbl} WHERE {q(missing['name'])} IS NULL"})
    return out
